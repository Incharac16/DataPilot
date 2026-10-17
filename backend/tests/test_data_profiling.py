import tempfile
import unittest
from pathlib import Path

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.services.chart_generator import generate_charts
from app.services.data_loader import DatasetStore, _clean_column_names, load_dataframe
from app.services.data_summary import build_dataset_summary, summarize_column
from app.services.data_types import semantic_type
from app.services.factual_analytics import calculate_dataset_facts
from app.services.insights_service import generate_grounded_insights
from app.config import settings
from unittest.mock import patch


class DataProfilingTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dataset_ids: list[str] = []
        self.data_dir_patcher = patch.object(settings, "data_dir", self.temp_dir.name)
        self.data_dir_patcher.start()

    def tearDown(self):
        self.temp_dir.cleanup()
        self.data_dir_patcher.stop()

    def register(self, filename: str, df: pd.DataFrame) -> str:
        path = Path(self.temp_dir.name) / filename
        if path.suffix == ".csv":
            df.to_csv(path, index=False)
        else:
            df.to_excel(path, index=False)
        dataset_id = path.stem
        DatasetStore.register(path, filename, dataset_id)
        self.dataset_ids.append(dataset_id)
        return dataset_id

    def test_numeric_column_has_measurement_statistics(self):
        column = summarize_column(pd.DataFrame({"Revenue": [10, 20, 30]}), "Revenue")

        self.assertEqual(column.semantic_type, "numeric")
        self.assertEqual(column.mean, 20)
        self.assertEqual(column.median, 20)
        self.assertEqual(column.min, 10)
        self.assertEqual(column.max, 30)
        self.assertEqual(column.top_value_status, "no_repeated_value")

    def test_unique_text_and_frequency_ties_are_explicit(self):
        unique = summarize_column(pd.DataFrame({"Name": ["Ada", "Lin", "Jo"]}), "Name")
        tied = summarize_column(pd.DataFrame({"Name": ["A", "A", "B", "B"]}), "Name")
        repeated = summarize_column(pd.DataFrame({"Name": ["A", "A", "B"]}), "Name")

        self.assertEqual(unique.semantic_type, "categorical")
        self.assertEqual(unique.top_value, "No repeated value")
        self.assertEqual(tied.top_value_status, "tie")
        self.assertEqual(tied.top_value_ties, ["A", "B"])
        self.assertEqual(tied.top_value_count, 2)
        self.assertEqual(repeated.top_value, "A")
        self.assertEqual(repeated.top_value_count, 2)

    def test_boolean_and_datetime_profiles_are_type_aware(self):
        df = pd.DataFrame(
            {
                "Active": pd.Series([True, False, True], dtype="boolean"),
                "Created At": pd.to_datetime(["2025-01-01", "2025-01-02", "2025-01-03"]),
            }
        )
        boolean = summarize_column(df, "Active")
        date = summarize_column(df, "Created At")

        self.assertEqual(boolean.semantic_type, "boolean")
        self.assertEqual((boolean.true_count, boolean.false_count), (2, 1))
        self.assertEqual((boolean.true_pct, boolean.false_pct), (66.67, 33.33))
        self.assertIsNone(boolean.mean)
        self.assertEqual(date.semantic_type, "datetime")
        self.assertEqual(date.earliest, "2025-01-01T00:00:00")
        self.assertEqual(date.latest, "2025-01-03T00:00:00")

    def test_identifiers_and_leading_zero_values_are_preserved(self):
        dataset_id = self.register(
            "identifiers.csv",
            pd.DataFrame(
                {
                    "Contact No": ["0012345678", "0012345678", "0098765432"],
                    "Email": ["ada@example.com", "ada@example.com", "lin@example.com"],
                    "Candidate ID": ["0007", "0008", "0009"],
                }
            ),
        )

        df = load_dataframe(dataset_id)
        summary = build_dataset_summary(dataset_id)
        by_name = {column.name: column for column in summary.column_summaries}

        self.assertEqual(df["Candidate ID"].tolist(), ["0007", "0008", "0009"])
        self.assertEqual(by_name["Contact No"].semantic_type, "identifier")
        self.assertIsNone(by_name["Contact No"].mean)
        self.assertEqual(by_name["Contact No"].duplicate_value_count, 1)
        self.assertEqual(by_name["Email"].semantic_type, "identifier")
        self.assertEqual(by_name["Candidate ID"].semantic_type, "identifier")

    def test_excel_numeric_identifier_formatting_preserves_leading_zeros(self):
        path = Path(self.temp_dir.name) / "formatted-ids.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Candidate ID"])
        for value in (7, 8, 9):
            worksheet.append([value])
            worksheet.cell(row=worksheet.max_row, column=1).number_format = "0000"
        workbook.save(path)
        dataset_id = path.stem
        DatasetStore.register(path, path.name, dataset_id)
        self.dataset_ids.append(dataset_id)

        df = load_dataframe(dataset_id)

        self.assertEqual(df["Candidate ID"].tolist(), ["0007", "0008", "0009"])

    def test_numeric_phone_ids_are_excluded_from_charts(self):
        dataset_id = self.register(
            "phone-chart.csv",
            pd.DataFrame({"Contact No": [1234567890, 1234567891], "Amount": [10, 20]}),
        )

        result = generate_charts(dataset_id)

        self.assertTrue(result.charts)
        self.assertTrue(all("Contact No" not in chart.title for chart in result.charts))
        self.assertTrue(all("Contact No" not in str(chart.figure) for chart in result.charts))
        self.assertEqual(
            semantic_type("Measurement", pd.Series([1234567890, 1234567891]))[0],
            "numeric",
        )

    def test_empty_values_whitespace_and_duplicate_counts_are_distinct(self):
        dataset_id = self.register(
            "missing-duplicates.csv",
            pd.DataFrame(
                {
                    "Candidate ID": ["1", "1", "2", None],
                    "Label": ["same", "same", " ", "other"],
                }
            ),
        )
        summary = build_dataset_summary(dataset_id)
        by_name = {column.name: column for column in summary.column_summaries}

        self.assertEqual(summary.duplicate_rows, 1)
        self.assertEqual(by_name["Candidate ID"].duplicate_value_count, 1)
        self.assertEqual(by_name["Candidate ID"].missing_count, 1)
        self.assertEqual(by_name["Label"].missing_count, 1)
        empty_column = summarize_column(pd.DataFrame({"Empty": [None, None]}), "Empty")
        self.assertEqual(empty_column.top_value, "No available values")
        self.assertEqual(empty_column.semantic_type, "unknown")
        self.assertIsNone(empty_column.mean)

    def test_named_iso_dates_are_parsed_but_ambiguous_dates_are_not(self):
        dataset_id = self.register(
            "date-inference.csv",
            pd.DataFrame(
                {
                    "Created At": ["2025-03-01", "2025-03-02"],
                    "Event Date": ["01/02/2025", "02/03/2025"],
                    "Legacy Date": ["20250301", "20250302"],
                    "Product Code": ["001", "002"],
                }
            ),
        )
        df = load_dataframe(dataset_id)

        self.assertTrue(pd.api.types.is_datetime64_any_dtype(df["Created At"]))
        self.assertFalse(pd.api.types.is_datetime64_any_dtype(df["Event Date"]))
        self.assertFalse(pd.api.types.is_numeric_dtype(df["Legacy Date"]))
        self.assertEqual(df["Product Code"].tolist(), ["001", "002"])

    def test_csv_missing_tokens_do_not_make_zero_or_false_missing(self):
        path = Path(self.temp_dir.name) / "missing-tokens.csv"
        path.write_text(
            'value,other\nNaN,x\nNone,x\nNaT,x\n"   ",x\n0,x\nFalse,x\n',
            encoding="utf-8",
        )
        dataset_id = path.stem
        DatasetStore.register(path, path.name, dataset_id)
        self.dataset_ids.append(dataset_id)

        df = load_dataframe(dataset_id)

        self.assertEqual(int(df["value"].isna().sum()), 4)
        self.assertIn("0", df["value"].dropna().tolist())
        self.assertIn("False", df["value"].dropna().tolist())

    def test_duplicate_excel_headers_are_unique_and_stable(self):
        df = _clean_column_names(pd.DataFrame([[1, 2, 3]], columns=["Value", "Value", " "]))

        self.assertEqual(list(df.columns), ["Value", "Value (2)", "Column 3"])

    def test_csv_and_excel_uploads_use_the_same_normalized_data(self):
        csv_bytes = b"Name,Amount\nAda,10\nLin,20\n"
        excel_path = Path(self.temp_dir.name) / "source.xlsx"
        pd.DataFrame({"Name": ["Ada", "Lin"], "Amount": [10, 20]}).to_excel(excel_path, index=False)
        excel_bytes = excel_path.read_bytes()

        with patch.object(settings, "upload_dir", self.temp_dir.name):
            client = TestClient(app)
            csv_response = client.post(
                "/api/upload",
                files={"file": ("dataset.csv", csv_bytes, "text/csv")},
            )
            excel_response = client.post(
                "/api/upload",
                files={
                    "file": (
                        "dataset.xlsx",
                        excel_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

        self.assertEqual(csv_response.status_code, 200)
        self.assertEqual(excel_response.status_code, 200)
        self.assertEqual(
            (csv_response.json()["rows"], csv_response.json()["columns"]),
            (excel_response.json()["rows"], excel_response.json()["columns"]),
        )
        self.dataset_ids.extend([csv_response.json()["id"], excel_response.json()["id"]])

    def test_data_quality_scores_follow_documented_formula(self):
        dataset_id = self.register(
            "quality.csv",
            pd.DataFrame(
                {
                    "Email": ["valid@example.com", "bad", "valid@example.com"],
                    "Amount": [10, None, 10],
                }
            ),
        )

        summary = build_dataset_summary(dataset_id)
        scores = {component.name: component.score for component in summary.data_quality.components}

        self.assertEqual(scores["Missing data"], 83.33)
        self.assertAlmostEqual(scores["Duplicate rows"], 66.67, places=2)
        self.assertEqual(scores["Type/format consistency"], 66.67)
        self.assertAlmostEqual(
            summary.data_quality.overall_score,
            sum(scores.values()) / 3,
            places=2,
        )
        self.assertTrue(summary.data_quality.recommendations)

    def test_potential_outliers_are_reported_without_affecting_score(self):
        dataset_id = self.register(
            "outliers.csv",
            pd.DataFrame({"Amount": [1, 2, 3, 4, 1000]}),
        )

        result = build_dataset_summary(dataset_id)

        self.assertEqual(result.data_quality.overall_score, 100)
        self.assertEqual(result.data_quality.potential_outliers[0].affected_count, 1)
        self.assertIn("1000", result.data_quality.potential_outliers[0].message)

    def test_identifier_only_chart_response_is_diagnostic(self):
        dataset_id = self.register(
            "ids-only.csv",
            pd.DataFrame({"Candidate ID": ["001", "002", "003"]}),
        )

        result = generate_charts(dataset_id)

        self.assertEqual(result.charts, [])
        self.assertTrue(
            any(
                "Eligible numeric measurement columns: []" in item
                for item in result.diagnostics
            )
        )
        unsupported_types = {chart.type for chart in result.unsupported_charts}
        self.assertEqual(
            unsupported_types,
            {
                "Bar chart",
                "Pie / donut chart",
                "Histogram",
                "Line chart",
                "Scatter plot",
                "Box plot",
                "Correlation heatmap",
            },
        )

    def test_categorical_dataset_reports_only_unsupported_chart_types(self):
        dataset_id = self.register(
            "categories-only.csv",
            pd.DataFrame(
                {
                    "Region": ["North", "South", "North", "West"],
                    "Product": ["A", "B", "A", "C"],
                }
            ),
        )

        result = generate_charts(dataset_id)

        self.assertTrue(all(chart.type == "pie" for chart in result.charts))
        unsupported = {chart.type: chart.reason for chart in result.unsupported_charts}
        self.assertEqual(
            set(unsupported),
            {
                "Bar chart",
                "Histogram",
                "Line chart",
                "Scatter plot",
                "Box plot",
                "Correlation heatmap",
            },
        )
        self.assertIn("no suitable numerical columns", unsupported["Histogram"])
        self.assertIn("date/time", unsupported["Line chart"])

    def test_numeric_only_dataset_gets_ungrouped_box_plot(self):
        dataset_id = self.register(
            "numeric-only.csv",
            pd.DataFrame({"Revenue": [10, 20, 30, 40]}),
        )

        result = generate_charts(dataset_id)

        box_charts = [chart for chart in result.charts if chart.type == "box"]
        self.assertEqual(len(box_charts), 1)
        self.assertEqual(box_charts[0].title, "Distribution of Revenue")
        self.assertIsNone(box_charts[0].x_column)
        self.assertEqual(box_charts[0].y_column, "Revenue")

    def test_wide_dataset_keeps_chart_and_heatmap_sizes_bounded(self):
        row_count = 30
        data = {
            f"measurement_{column}_with_a_long_column_name": [
                row + column for row in range(row_count)
            ]
            for column in range(20)
        }
        data.update(
            {
                f"category_{column}_with_a_long_column_name": [
                    f"category-{(row + column) % 3}" for row in range(row_count)
                ]
                for column in range(20)
            }
        )
        dataset_id = self.register("wide-dataset.csv", pd.DataFrame(data))

        result = generate_charts(dataset_id)

        self.assertLessEqual(len(result.charts), 6)
        heatmap = next(chart for chart in result.charts if chart.type == "heatmap")
        self.assertEqual(heatmap.figure["data"][0]["z"]["shape"], "15, 15")
        self.assertTrue(
            any("limited to the first 15" in message for message in result.diagnostics)
        )

    def test_chart_processing_error_is_returned_as_server_error(self):
        dataset_id = self.register(
            "chart-processing-error.csv",
            pd.DataFrame({"Revenue": [10, 20, 30]}),
        )

        with patch(
            "app.services.chart_generator._build_chart",
            side_effect=RuntimeError("Plotly processing failed"),
        ):
            response = TestClient(app, raise_server_exceptions=False).get(
                f"/api/datasets/{dataset_id}/charts"
            )

        self.assertEqual(response.status_code, 500)
        self.assertIn("Plotly processing failed", response.json()["detail"])

    def test_ai_facts_are_calculated_and_returned_separately_from_model_text(self):
        dataset_id = self.register(
            "grounded-sales.csv",
            pd.DataFrame(
                {
                    "Region": ["North", "South", "North", "South"],
                    "Revenue": [100, 80, 50, 70],
                    "Profit": [10, 20, 30, 40],
                    "Product_Category": ["A", "B", "A", "B"],
                    "Notes": [None, "ok", None, "ok"],
                }
            ),
        )

        class MockProvider:
            prompt = ""

            def generate(self, system_prompt, user_prompt):
                self.prompt = user_prompt
                return "The actual calculations are listed separately."

        provider = MockProvider()
        question = "Which region has the highest total revenue?"
        with patch("app.services.insights_service.get_ai_provider", return_value=provider):
            answer, facts = generate_grounded_insights(dataset_id, question)

        self.assertEqual(answer, "The actual calculations are listed separately.")
        self.assertIn("highest is North (150)", "\n".join(facts))
        self.assertIn('"total": 150', provider.prompt)
        self.assertNotIn("North, South, North, South", provider.prompt)

        calculated = calculate_dataset_facts(
            load_dataframe(dataset_id),
            "What is the average profit and category shares for Product_Category?",
        )
        self.assertEqual(calculated["columns"]["Profit"]["statistics"]["mean"], 25.0)
        self.assertEqual(
            [item["percentage_of_non_missing"] for item in calculated["columns"]["Product_Category"]["category_counts"]],
            [50.0, 50.0],
        )
        self.assertEqual(calculated["overview"]["missing_by_column"]["Notes"], 2)

    def test_ai_provider_failures_are_not_returned_as_successful_insights(self):
        dataset_id = self.register(
            "provider-failure.csv",
            pd.DataFrame({"Revenue": [10, 20, 30]}),
        )
        with patch(
            "app.services.insights_service.get_ai_provider",
            side_effect=TimeoutError("provider timeout"),
        ):
            with self.assertRaisesRegex(RuntimeError, "could not complete"):
                generate_grounded_insights(dataset_id, "Summarize revenue")

    def test_pdf_report_contains_actual_dataset_values_and_download_headers(self):
        dataset_id = self.register(
            "sales-report.csv",
            pd.DataFrame(
                {
                    "Region": ["North", "South", "North"],
                    "Revenue": [100, 50, None],
                    "Profit": [20, 10, 40],
                }
            ),
        )

        response = TestClient(app).get(f"/api/datasets/{dataset_id}/report.pdf")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertEqual(
            response.headers["content-disposition"],
            'attachment; filename="sales-report-report.pdf"',
        )
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(b"sales-report.csv", response.content)
        self.assertIn(b"Generated:", response.content)
        self.assertIn(b"Rows: 3", response.content)
        self.assertIn(b"Columns: 3", response.content)
        self.assertIn(b"Data quality and reliability", response.content)
        self.assertIn(b"Suggested action", response.content)
        self.assertIn(b"Bar Chart", response.content)
        self.assertIn(b"Revenue by Region", response.content)
        self.assertIn(b"Histogram", response.content)
        self.assertIn(b"Distribution of Revenue", response.content)
        self.assertIn(b"Average 75", response.content)
        self.assertIn(b"Missing values", response.content)
        self.assertIn(b"AI-generated insights", response.content)
        self.assertIn(b"No saved AI-generated insight", response.content)
        self.assertNotIn(b"Example repeated rows", response.content)

    def test_pdf_report_includes_every_analytics_chart_with_both_labels(self):
        dataset_id = self.register(
            "chart-report.csv",
            pd.DataFrame(
                {
                    "Order_Date": pd.date_range("2025-01-01", periods=4),
                    "Product_Category": ["A", "B", "A", "C"],
                    "Units_Sold": [10, 20, 30, 50],
                    "Unit_Price": [4, 5, 7, 9],
                }
            ),
        )
        client = TestClient(app)

        charts_response = client.get(f"/api/datasets/{dataset_id}/charts")
        report_response = client.get(f"/api/datasets/{dataset_id}/report.pdf")

        self.assertEqual(charts_response.status_code, 200)
        self.assertEqual(report_response.status_code, 200)
        chart_labels = {
            "histogram": "Histogram",
            "bar": "Bar Chart",
            "line": "Line Chart",
            "scatter": "Scatter Plot",
            "box": "Box Plot",
            "heatmap": "Correlation Heatmap",
            "pie": "Pie / Donut Chart",
        }
        for chart in charts_response.json()["charts"]:
            self.assertIn(chart_labels[chart["type"]].encode(), report_response.content)
            self.assertIn(chart["title"].encode(), report_response.content)
        self.assertIn(b"Shows pairwise correlation values", report_response.content)
        self.assertIn(b"without asserting a cause", report_response.content)

    def test_pdf_report_paginates_long_column_profiles(self):
        frame = {
            f"long_column_{index}_" + ("x" * 48): ["first", "second", "third"]
            for index in range(30)
        }
        dataset_id = self.register("many-columns.csv", pd.DataFrame(frame))

        response = TestClient(app).get(f"/api/datasets/{dataset_id}/report.pdf")
        page_count = response.content.count(b"/Type /Page\n")

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreaterEqual(page_count, 3)
        self.assertIn(b"long_column_0_", response.content)
        self.assertIn(b"Dataset profile and column statistics", response.content)
        self.assertIn(b"Generated by DataPilot", response.content)

    def test_pdf_report_handles_a_smaller_dataset_with_a_pie_chart(self):
        dataset_id = self.register(
            "small-pie-report.csv",
            pd.DataFrame(
                {
                    "Region": ["North", "South", "East", "West"],
                    "Revenue": [10, 30, 20, 40],
                }
            ),
        )
        charts = generate_charts(dataset_id).charts
        response = TestClient(app).get(f"/api/datasets/{dataset_id}/report.pdf")

        self.assertLess(len(charts), 6)
        pie_chart = next(chart for chart in charts if chart.type == "pie")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Pie / Donut Chart", response.content)
        self.assertIn(pie_chart.title.encode(), response.content)
        self.assertIn(b" c ", response.content)

    def test_pdf_report_still_succeeds_when_optional_chart_generation_fails(self):
        dataset_id = self.register(
            "chart-failure-report.csv",
            pd.DataFrame({"Region": ["North", "South"], "Revenue": [100, 50]}),
        )

        with patch(
            "app.services.report_service._chart_drawings",
            side_effect=RuntimeError("optional chart failed"),
        ):
            response = TestClient(app).get(
                f"/api/datasets/{dataset_id}/report.pdf"
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(b"chart-failure-report.csv", response.content)
        self.assertIn(b"No charts were available", response.content)

    def test_pdf_report_handles_empty_dataset_and_missing_ids(self):
        empty_dataset_id = self.register(
            "empty-report.csv",
            pd.DataFrame({"Revenue": pd.Series(dtype="float64")}),
        )
        client = TestClient(app)

        empty_report = client.get(
            f"/api/datasets/{empty_dataset_id}/report.pdf"
        )
        missing_report = client.get(
            "/api/datasets/not-a-dataset/report.pdf"
        )

        self.assertEqual(empty_report.status_code, 200)
        self.assertTrue(empty_report.content.startswith(b"%PDF-"))
        self.assertIn(b"Rows: 0", empty_report.content)
        self.assertEqual(missing_report.status_code, 404)
        self.assertIn("not found", missing_report.json()["detail"].lower())

    def test_pdf_generation_error_is_returned_with_actionable_detail(self):
        dataset_id = self.register(
            "failed-report.csv",
            pd.DataFrame({"Revenue": [100, 50]}),
        )

        with patch(
            "app.routers.analysis.generate_pdf_report",
            side_effect=RuntimeError("report layout failed"),
        ):
            response = TestClient(app, raise_server_exceptions=False).get(
                f"/api/datasets/{dataset_id}/report.pdf"
            )

        self.assertEqual(response.status_code, 500)
        self.assertIn("PDF report generation failed", response.json()["detail"])
        self.assertNotIn("report layout failed", response.json()["detail"])

    def test_upload_to_analytics_insights_and_pdf_workflow(self):
        client = TestClient(app)
        rows = [
            "Order_Date,Region,Product_Category,Revenue,Profit,Units_Sold,Customer_Rating"
        ]
        for index in range(12):
            rows.append(
                f"2025-01-{index + 1:02d},{'North' if index % 2 == 0 else 'South'},"
                f"{'Hardware' if index % 3 else 'Software'},{100 + index * 10},"
                f"{20 + index * 2},{index + 1},{4 + index / 10:.1f}"
            )
        with patch.object(settings, "upload_dir", self.temp_dir.name):
            uploaded = client.post(
                "/api/upload",
                files={"file": ("workflow-sales.csv", ("\n".join(rows) + "\n").encode(), "text/csv")},
            )
        self.assertEqual(uploaded.status_code, 200)
        dataset_id = uploaded.json()["id"]

        preview = client.get(
            f"/api/datasets/{dataset_id}/summary",
            params={"page": 2, "page_size": 10, "search": "SOUTH"},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["total_rows"], 12)
        self.assertEqual(preview.json()["matching_rows"], 6)
        self.assertEqual(preview.json()["current_page"], 1)
        self.assertTrue(preview.json()["preview_rows"][0]["Order_Date"].startswith("2025-01-"))

        summary = client.get(f"/api/datasets/{dataset_id}/summary")
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary.json()["columns"], 7)
        self.assertEqual(summary.json()["missing_values_total"], 0)

        charts = client.get(f"/api/datasets/{dataset_id}/charts")
        self.assertEqual(charts.status_code, 200)
        self.assertLessEqual(len(charts.json()["charts"]), 6)
        self.assertTrue(
            any(chart["title"] == "Revenue by Product_Category" for chart in charts.json()["charts"])
        )
        self.assertTrue(any(chart["type"] == "heatmap" for chart in charts.json()["charts"]))

        class MockProvider:
            def generate(self, system_prompt, user_prompt):
                self.prompt = user_prompt
                return "Revenue trends upward in the provided records."

        provider = MockProvider()
        with patch("app.services.insights_service.get_ai_provider", return_value=provider):
            insights = client.post(
                f"/api/datasets/{dataset_id}/insights",
                json={"question": "What is the average revenue?"},
            )
        self.assertEqual(insights.status_code, 200)
        self.assertTrue(any("Revenue mean" in fact for fact in insights.json()["verified_facts"]))
        self.assertIn('"mean": 155.0', provider.prompt)

        report = client.get(f"/api/datasets/{dataset_id}/report.pdf")
        self.assertEqual(report.status_code, 200)
        self.assertTrue(report.content.startswith(b"%PDF-"))
        self.assertIn(b"workflow-sales.csv", report.content)

    def test_sales_dataset_generates_expected_chart_combinations(self):
        dataset_id = self.register(
            "sales-shortlist.csv",
            pd.DataFrame(
                {
                    "Order_Date": pd.date_range("2025-01-01", periods=8, freq="D"),
                    "Region": ["North", "South"] * 4,
                    "Product_Category": ["Hardware", "Software"] * 4,
                    "Revenue": [100, 130, 125, 155, 140, 180, 175, 200],
                    "Profit": [20, 25, 22, 31, 27, 38, 35, 44],
                    "Units_Sold": [2, 3, 2, 4, 3, 5, 4, 6],
                    "Customer_Rating": [4.1, 4.2, 4.0, 4.5, 4.3, 4.7, 4.6, 4.8],
                }
            ),
        )

        result = generate_charts(dataset_id)
        specifications = {(chart.type, chart.title) for chart in result.charts}

        self.assertLessEqual(len(result.charts), 6)
        self.assertIn(("bar", "Revenue by Product_Category"), specifications)
        self.assertIn(("histogram", "Distribution of Revenue"), specifications)
        self.assertIn(("line", "Revenue over Order_Date"), specifications)
        self.assertIn(("scatter", "Profit vs Revenue"), specifications)
        self.assertIn(("box", "Profit by Product_Category"), specifications)
        heatmap = next(chart for chart in result.charts if chart.type == "heatmap")
        self.assertEqual(heatmap.title, "Correlation Between Numeric Variables")
        line = next(chart for chart in result.charts if chart.type == "line")
        self.assertEqual(line.x_column, "Order_Date")
        self.assertEqual(line.y_column, "Revenue")
        self.assertTrue(all(chart.title for chart in result.charts))

    def test_chart_titles_reflect_the_current_dataset_columns(self):
        dataset_id = self.register(
            "different-columns.csv",
            pd.DataFrame(
                {
                    "Sales_Channel": ["Online", "Retail", "Online", "Retail"],
                    "Visits": [10, 15, 30, 45],
                    "Order_Date": pd.date_range("2025-01-01", periods=4),
                }
            ),
        )

        result = generate_charts(dataset_id)
        chart_titles = {chart.title for chart in result.charts}

        self.assertIn("Visits by Sales_Channel", chart_titles)
        self.assertIn("Visits over Order_Date", chart_titles)
        self.assertNotIn("Revenue by Product_Category", chart_titles)

    def test_chart_api_generates_selected_real_dataset_chart(self):
        dataset_id = self.register(
            "selected-chart.csv",
            pd.DataFrame(
                {
                    "Region": ["North", "South", "North", "South"],
                    "Revenue": [10, 20, 15, 25],
                }
            ),
        )

        response = TestClient(app).get(
            f"/api/datasets/{dataset_id}/charts",
            params={
                "chart_type": "bar",
                "x_column": "Region",
                "y_column": "Revenue",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["charts"][0]["title"], "Revenue by Region")
        self.assertEqual(payload["charts"][0]["x_column"], "Region")
        self.assertEqual(payload["charts"][0]["y_column"], "Revenue")
        self.assertTrue(payload["charts"][0]["figure"]["data"])
        self.assertIn("Revenue", payload["available_columns"]["numeric"])


if __name__ == "__main__":
    unittest.main()
