import json
import re
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
from openpyxl import load_workbook

from app.config import settings
from app.services.data_types import (
    named_date_column,
    named_identifier_category,
    normalize_dataframe_types,
)
from app.utils.file_validator import validate_file_size


def _clean_column_names(df: pd.DataFrame) -> pd.DataFrame:
    names: list[str] = []
    used: set[str] = set()

    for index, column in enumerate(df.columns, start=1):
        name = str(column).strip()
        if not name or name.lower().startswith("unnamed:"):
            name = f"Column {index}"

        unique_name = name
        suffix = 2
        while unique_name in used:
            unique_name = f"{name} ({suffix})"
            suffix += 1

        names.append(unique_name)
        used.add(unique_name)

    df.columns = names
    return df


def _find_excel_header_row(raw_df: pd.DataFrame, search_rows: int = 20) -> int | None:
    candidates: list[tuple[tuple[float, float, int, int], int]] = []

    for index, row in raw_df.iloc[:search_rows].iterrows():
        values = row.dropna().tolist()
        if len(values) < 2:
            continue

        text_values = [value for value in values if isinstance(value, str) and value.strip()]
        text_ratio = len(text_values) / len(values)
        unique_ratio = len({str(value).strip().casefold() for value in values}) / len(values)
        score = (text_ratio, unique_ratio, len(values), -int(index))
        candidates.append((score, int(index)))

    return max(candidates, default=(None, None))[1]


def _restore_formatted_identifier_values(
    path: Path,
    df: pd.DataFrame,
    header_row: int,
) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        for column_index, column_name in enumerate(df.columns, start=1):
            if not named_identifier_category(str(column_name))[0]:
                continue
            values = df.iloc[:, column_index - 1].astype(object).to_numpy(copy=True)
            cells = worksheet.iter_rows(
                min_row=header_row + 2,
                max_row=header_row + 1 + len(df),
                min_col=column_index,
                max_col=column_index,
            )
            for dataframe_row, (cell,) in enumerate(cells):
                if not isinstance(cell.value, (int, float)) or not re.fullmatch(
                    r"0+",
                    cell.number_format,
                ):
                    continue
                width = len(cell.number_format)
                values[dataframe_row] = f"{int(cell.value):0{width}d}"
            df.isetitem(column_index - 1, values)
    finally:
        workbook.close()

    return df


def load_dataframe_from_path(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        headers = pd.read_csv(path, nrows=0, keep_default_na=False).columns
        identifier_columns = {
            str(column): "string"
            for column in headers
            if named_identifier_category(str(column))[0] or named_date_column(str(column))
        }
        df = pd.read_csv(
            path,
            dtype=identifier_columns,
            keep_default_na=False,
        )
    elif path.suffix.lower() in {".xlsx", ".xls"}:
        raw_df = pd.read_excel(path, header=None)
        header_row = _find_excel_header_row(raw_df)
        actual_header_row = header_row if header_row is not None else 0
        df = (
            pd.read_excel(path, header=header_row)
            if header_row is not None
            else pd.read_excel(path)
        )
        if path.suffix.lower() == ".xlsx":
            df = _restore_formatted_identifier_values(
                path,
                df,
                actual_header_row,
            )
        df = df.dropna(axis=0, how="all")
    else:
        raise ValueError("Unsupported file type")

    df = _clean_column_names(df)
    for column in df.columns:
        if pd.api.types.is_object_dtype(df[column]) or pd.api.types.is_string_dtype(df[column]):
            values = df[column].map(
                lambda value: value.strip() if isinstance(value, str) else value
            )
            df[column] = values.map(
                lambda value: pd.NA
                if isinstance(value, str) and value.casefold() in {"", "nan", "none", "nat"}
                else value
            )

    return normalize_dataframe_types(df)


class DatasetStore:
    @classmethod
    @contextmanager
    def _connect(cls):
        database_path = Path(settings.data_dir) / "datasets.sqlite3"
        database_path.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(database_path)
        connection.row_factory = sqlite3.Row
        try:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS datasets (
                    id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    uploaded_at TEXT NOT NULL,
                    file_format TEXT NOT NULL,
                    row_count INTEGER NOT NULL,
                    column_count INTEGER NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    storage_path TEXT NOT NULL,
                    status TEXT NOT NULL
                )
                """
            )
            yield connection
            connection.commit()
        finally:
            connection.close()

    @classmethod
    def register(
        cls,
        file_path: Path,
        filename: str,
        dataset_id: str,
        row_count: int | None = None,
        column_count: int | None = None,
        size_bytes: int | None = None,
        uploaded_at: str | None = None,
    ) -> None:
        if row_count is None or column_count is None:
            df = load_dataframe_from_path(file_path)
            row_count = len(df)
            column_count = len(df.columns)
        metadata = (
            dataset_id,
            filename,
            uploaded_at or datetime.now(timezone.utc).isoformat(),
            file_path.suffix.lower().lstrip("."),
            row_count,
            column_count,
            size_bytes if size_bytes is not None else file_path.stat().st_size,
            str(file_path.resolve()),
            "ready",
        )
        with cls._connect() as connection:
            connection.execute(
                """
                INSERT OR REPLACE INTO datasets (
                    id, filename, uploaded_at, file_format, row_count,
                    column_count, size_bytes, storage_path, status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                metadata,
            )

    @classmethod
    def get(cls, dataset_id: str) -> dict[str, Any]:
        with cls._connect() as connection:
            row = connection.execute(
                "SELECT * FROM datasets WHERE id = ?", (dataset_id,)
            ).fetchone()
        if row is None:
            raise KeyError(f"Dataset '{dataset_id}' not found")
        metadata = dict(row)
        metadata["path"] = Path(metadata.pop("storage_path"))
        if not metadata["path"].is_file():
            raise FileNotFoundError(
                f"Stored file for dataset '{dataset_id}' is missing."
            )
        return metadata

    @classmethod
    def exists(cls, dataset_id: str) -> bool:
        with cls._connect() as connection:
            row = connection.execute(
                "SELECT 1 FROM datasets WHERE id = ?", (dataset_id,)
            ).fetchone()
        return row is not None

    @classmethod
    def list(cls) -> list[dict[str, Any]]:
        with cls._connect() as connection:
            rows = connection.execute(
                """
                SELECT id, filename, uploaded_at, file_format, row_count,
                       column_count, size_bytes, storage_path, status
                FROM datasets ORDER BY uploaded_at DESC
                """
            ).fetchall()
        records = []
        for row in rows:
            record = dict(row)
            record["status"] = (
                "ready" if Path(record.pop("storage_path")).is_file() else "missing"
            )
            records.append(record)
        return records

    @classmethod
    def delete(cls, dataset_id: str) -> None:
        with cls._connect() as connection:
            row = connection.execute(
                "SELECT storage_path FROM datasets WHERE id = ?", (dataset_id,)
            ).fetchone()
            if row is None:
                raise KeyError(f"Dataset '{dataset_id}' not found")
            stored_path = Path(row["storage_path"]).resolve()
            upload_root = Path(settings.upload_dir).resolve()
            if stored_path.parent != upload_root:
                raise ValueError("Stored file is outside the configured upload directory.")
            stored_path.unlink(missing_ok=True)
            connection.execute("DELETE FROM datasets WHERE id = ?", (dataset_id,))


def save_upload(file_content: BinaryIO, filename: str, max_upload_mb: int):
    upload_dir = Path(settings.upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)

    dataset_id = str(uuid.uuid4())
    ext = Path(filename).suffix.lower()
    file_path = upload_dir / f"{dataset_id}{ext}"

    total_size = 0
    file_created = False
    upload_complete = False
    try:
        with file_path.open("xb") as destination:
            file_created = True
            while chunk := file_content.read(1024 * 1024):
                total_size += len(chunk)
                validate_file_size(total_size, max_upload_mb)
                destination.write(chunk)

        if total_size == 0:
            raise ValueError("The uploaded file is empty.")
        upload_complete = True
    finally:
        if file_created and not upload_complete:
            file_path.unlink(missing_ok=True)

    return dataset_id, file_path



def load_dataframe(dataset_id: str):
    meta = DatasetStore.get(dataset_id)
    return load_dataframe_from_path(Path(meta["path"]))



def dataframe_preview(df: pd.DataFrame):
    preview = df.copy()
    preview = preview.where(preview.notna(), None)
    return json.loads(preview.to_json(orient="records", date_format="iso"))